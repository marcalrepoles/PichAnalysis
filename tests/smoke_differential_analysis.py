"""Offline preparation -> limma -> historical differential-run smoke."""
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from pichanalysis.core.differential_analysis import (
    DifferentialParameters, PLOTS, TABLES, load_run, run_differential_analysis)
from pichanalysis.core.differential_preparation import (
    PreparationParameters, run_differential_preparation)
from pichanalysis.core.r_runtime import RRuntime
from smoke_differential_preparation import fixture_project


def smoke():
    with tempfile.TemporaryDirectory() as folder:
        project, columns, source = fixture_project(Path(folder))
        data = pd.read_csv(source)
        data.loc[data.ProteinGroup == "P11", ["B1", "B2", "B3"]] = [20, 21, 22]
        data.loc[data.ProteinGroup == "P12", ["B1", "B2", "B3"]] = [55, 56, 57]
        data.loc[data.ProteinGroup == "P13", columns] = [100, 101, 102, 10, 11, 12]
        data.to_csv(source, index=False)
        runtime = RRuntime()
        run_x = run_differential_preparation(project, runtime, run_id="prep_x",
            parameters=PreparationParameters("A", "B", tuple(columns)))
        run_y = run_differential_preparation(project, runtime, run_id="prep_y",
            parameters=PreparationParameters("A", "B", tuple(columns),
                normalization="median_center", imputation_method="MinProb"))
        a = run_differential_analysis(project, runtime, run_id="stat_a",
            parameters=DifferentialParameters("prep_x"))
        assert set(a["tables"]) == set(TABLES)
        assert len(a["plots"]) == len(PLOTS)
        assert a["workbook"].is_file()
        workbook = load_workbook(a["workbook"], read_only=True)
        assert {"Summary", "Preparation run", "Design", "All results", "Tested results",
            "FDR significant", "Significant", "Not tested", "Qualitative candidates",
            "Model audit", "eBayes audit"}.issubset(workbook.sheetnames)
        workbook.close()
        assert a["metadata"]["comparison_direction"] == "Condition A - Condition B"
        assert a["metadata"]["prepared_scale"] == "log2"
        results = a["tables"]["all_results"]
        assert len(results) == len(run_x["tables"]["prepared_matrix"])
        assert results.loc[0, "feature_id"] != results.loc[1, "feature_id"]
        assert results.loc[2, "display_identifier"] == "P3;P4"
        assert "ROW000005" not in results.feature_id.tolist()
        assert "ROW000005" in a["tables"]["qualitative_candidates"].feature_id.tolist()
        assert set(a["tables"]["qualitative_candidates"].columns).isdisjoint({"P.Value", "log2FC", "fold_change"})
        assert len(set(results.df_residual)) > 1
        assert results.loc[results.feature_id == "ROW000011", "effect"].astype(float).iloc[0] > 0
        assert results.loc[results.feature_id == "ROW000001", "effect"].astype(float).iloc[0] < 0
        assert results.loc[results.feature_id == "ROW000004", "n_A_model"].iloc[0] == "2"
        assert (results.loc[results.tested == "TRUE", "adj.P.Val"].astype(float) >= 0).all()
        assert (results.imputed_cell_count == "0").all()
        b = run_differential_analysis(project, runtime, run_id="stat_b",
            parameters=DifferentialParameters("prep_y", ebayes_trend=True, ebayes_robust=True,
                minimum_absolute_effect=1))
        assert b["metadata"]["parent_preparation_run_id"] == "prep_y"
        assert b["metadata"]["ebayes_trend"] is True
        assert b["metadata"]["ebayes_robust"] is True
        assert (b["run_root"] / "plots/imputation_diagnostic.png").is_file()
        assert (b["run_root"] / "plots/imputation_diagnostic.pdf").is_file()
        assert (b["tables"]["all_results"].contains_imputed_values == "TRUE").any()
        source.write_text("ProteinGroup,A1,A2,A3,B1,B2,B3\nChanged,1,1,1,1,1,1\n", encoding="utf-8")
        old_results = results.copy()
        old = load_run(project, "stat_a")
        assert old["tables"]["all_results"].equals(old_results)
        assert old["metadata"]["parent_preparation_run_id"] == "prep_x"
        hidden = run_x["run_root"].with_name("prep_x_hidden")
        run_x["run_root"].rename(hidden)
        try:
            assert load_run(project, "stat_a")["tables"]["all_results"].equals(old_results)
        finally:
            hidden.rename(run_x["run_root"])
        audit = a["tables"]["ebayes_audit"].iloc[0]
        print("R version:", audit.r_version, "limma version:", audit.limma_version)
        print("Features tested:", audit.features_tested,
            "residual df:", audit.residual_df_min, audit.residual_df_max,
            "df prior:", audit.df_prior_median,
            "significant:", len(a["tables"]["significant_results"]))
        print("Differential statistics offline smoke passed")


if __name__ == "__main__":
    smoke()
